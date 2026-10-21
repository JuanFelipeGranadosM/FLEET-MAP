"""Carga BMC con un libro sintético (los libros reales no se versionan).

Con FLEETMAP_TEST_DSN (PostgreSQL + PostGIS vacío) también aplica los esquemas y revisa las vistas.
"""
import os
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
import pytest

from app.domain.haulage import StandbyRule, m3km, standby_by_slot
from app.ingest.bmc import classify, load_all, summary
from app.ingest.store import to_sql

DB = Path(__file__).resolve().parents[1] / "db"


def _f12(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Acta"
    ws["B2"] = "ACTA DE CORTE No. 01"
    ws["B5"], ws["D5"], ws["E5"], ws["G5"] = "Contrato No.:", "CT-TEST", "CC-SC:", "3.02.04"
    ws["H5"], ws["J5"], ws["L5"] = "Fecha Elaboración Acta:", datetime(2026, 9, 21), "Contrato Inicial"
    ws["P5"], ws["Q5"], ws["R5"], ws["S5"] = 4, datetime(2026, 8, 20), datetime(2026, 12, 18), 1_000_000
    ws["B6"], ws["D6"], ws["F6"], ws["G6"] = "Contratante:", "CLIENTE SA", "Contratista:", "ORTIZ CONSTRUCCIONES"
    ws["B7"], ws["D7"], ws["F7"], ws["G7"] = "Periodo Facturado:", datetime(2026, 8, 20), "a:", datetime(2026, 9, 16)
    for col, v in zip("BCEFGHJKLO", ["3.02.04", "Transporte de mezcla", "m3-km", 1000, 900, 900000, 1000, 1300, 1300000, 4740]):
        ws[f"{col}13"] = v

    d = wb.create_sheet("Detalle Acta 01")
    d.append(["Fecha", "Recibo", "Placa", "Volumen suelto (m3)", "Conductor", "Proveedor", "Ruta", "Abs Descargue",
              "Origen (Planta)", "Destino (Frente)", "Distancia Acarreo (km)", "Total (m3-km) Vol Suelto"])
    d.append([datetime(2026, 9, 1), "1", "AAA001", 20, "Op A", "G. Ortiz", "R1", 100, "Planta", "Frente", 97.2, 2914, None, None, None, "PARCHEO"])
    d.append([datetime(2026, 9, 1), "2", "AAA002", 20, "Op B", "G. Ortiz", "R1", 100, "Planta", "Frente", 97.2, 1944])
    d.append([datetime(2026, 9, 1), "3", "ZZZ999", 14, "Op C", "Tercero SAS", "R1", 100, "Planta", "Frente", 97.2, 1361])

    s = wb.create_sheet("STAND BY")
    s.append(["TRACTOCAMIÓN", "PLACA", "m3", "Km", "Viajes", "SB MES COMPLETO", "INICIO", "FIN", "DÍAS", "SB RELATIVO"])
    s.append(["T001", "AAA001", 20, 94, 15, 28200, datetime(2026, 8, 20), datetime(2026, 9, 15), 27, 25380])
    s.append(["T002", "AAA002", 20, 94, 15, 28200, datetime(2026, 8, 20), datetime(2026, 9, 15), 27, 25380])
    wb.save(path)


def _panel(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registros"
    ws.append(["FECHA", "CODIGO", "PLACA", "CONDUCTOR", "KM INI", "KM FIN", "TOTAL KMS", "HRS IN OP", "HRS FIN OP",
               "TOTAL HRS T", "M3 TRANS", "CANT.VIAJES", "TRAYECTO", "OBS"])
    ws.append([datetime(2026, 9, 1), None, "AAA001", "Op A", 1000, 1200, None, time(23, 0), time(16, 0), None, 20, 1,
               "Planta - Frente (acarreo 97.2 km)", "Parte 16677. Jornada nocturna"])
    ws.append([datetime(2026, 9, 2), None, "AAA001", "Op A", 1200, 1200, None, time(7, 0), time(16, 0), None, 0, 0,
               "Movimiento Interno Taller", None])
    wb.save(path)


@pytest.fixture
def load(tmp_path):
    _f12(tmp_path / "f12.xlsx")
    _panel(tmp_path / "panel.xlsx")
    return load_all(tmp_path / "f12.xlsx", [], tmp_path / "panel.xlsx", None, None)


def test_rules():
    assert m3km(24.03, 66.5) == 1598 and m3km(22, 145.7) == 3205
    rule = StandbyRule(20, 94, 15)
    assert rule.month_m3km == 28200 and rule.minimum(27) == 25380
    # dos equipos que se relevan suman sus días en un solo cupo
    res = standby_by_slot(rule, [{"plate": "A", "slot": "S", "days_available": 16}, {"plate": "B", "slot": "S", "days_available": 13}],
                          [{"plate": "A", "m3": 20, "haul_km": 97.2}, {"plate": "B", "m3": 20, "haul_km": 96.8}])
    assert len(res) == 1 and res[0].minimum == 27260 and res[0].executed == 1944 + 1936
    assert classify("No programado", 1) == "trabajando" and classify("Disponible en planta", 0) == "disponible"
    assert classify("Equipo en Mantenimiento Taller", 0) == "taller" and classify(None, None) == "sin_registro"


def test_parse(load):
    c = load.contract
    assert (c.id, c.client, c.start, c.end, c.value) == ("CT-TEST", "CLIENTE SA", date(2026, 8, 20), date(2026, 12, 18), 1_000_000)
    it = load.items[0]
    assert (it.code, it.unit, it.price, it.price_updated) == ("3.02.04", "m3-km", 900, 1300)
    assert load.rule == StandbyRule(20, 94, 15)
    s = summary(load)[0]
    # el m³-km pegado como constante (2914 en vez de 20 × 97,2 = 1944) aparece como diferencia
    assert (s["viajes"], s["m3km_acta"], s["m3km_calculado"], s["diferencia"], s["terceros_m3km"]) == (2, 4858, 3888, 970, 1361)
    assert s["standby"] == {"days": 54, "minimum": 50760, "shortfall": 50760 - 3888}
    r = load.reports[0]
    assert (r.report_no, r.status, r.origin, r.destination) == ("16677", "trabajando", "Planta", "Frente")
    assert load.reports[1].status == "taller"


@pytest.mark.skipif(not os.environ.get("FLEETMAP_TEST_DSN"), reason="sin FLEETMAP_TEST_DSN")
def test_postgres(load):
    psycopg = pytest.importorskip("psycopg")
    with psycopg.connect(os.environ["FLEETMAP_TEST_DSN"], autocommit=True) as conn:
        conn.execute("DROP SCHEMA IF EXISTS public CASCADE; CREATE SCHEMA public;")
        for f in ("schema.sql", "002_operacion_facturacion.sql"):
            conn.execute((DB / f).read_text(encoding="utf-8"))
        sql = to_sql(load)
        conn.execute(sql)
        conn.execute(sql)  # idempotente
        rec = conn.execute("SELECT trips, m3km_billed, m3km_calc FROM v_cut_reconciliation").fetchone()
        assert tuple(map(float, rec)) == (2, 4858, 3888)
        assert float(conn.execute("SELECT tariff_at('CT-TEST', '3.02.04', '2026-09-01')").fetchone()[0]) == 900
        h, km = conn.execute("SELECT hours, km FROM daily_report WHERE day = '2026-09-01'").fetchone()
        assert (float(h), float(km)) == (17, 200)   # 23:00 → 16:00 del día siguiente
        assert conn.execute("SELECT count(*) FROM haul_ticket").fetchone()[0] == 3
