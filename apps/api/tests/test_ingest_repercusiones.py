"""Carga del histórico de repercusiones con un libro sintético."""
import os
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from app.ingest.repercusiones import parse, rate_history, short_code, summary, to_sql

DB = Path(__file__).resolve().parents[1] / "db"
HDR = ["CORTE", "TIPO DE COSTE", "TIPO DE EQUIPO", "CÓDIGO EQUIPO", "HORAS ", "%", "TOTAL HORAS TRABAJADAS Y/O DISPONIBLES",
       "TARIFA HORA", "TOTAL REPERCUSIÓN", "NOMBRE PROYECTO", "PROYECTO DESTINO", "CENTRO DE COSTO DESTINO", "FAMILIA",
       "NOMBRE FAMILIA", "PROYECTO", "TIPO DE INGRESO"]


@pytest.fixture
def book(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HISTORICO"
    ws.append(["ACTA DE REPERCUSION MAQUINARIA"])
    ws.append(HDR)
    ws.append([datetime(2026, 1, 1), "VENTA", "TRACTOCAMIÓN X", "40BUP0TC001", None, None, 100, 1000, 100000, "T2 - UF1", "3701", "3701.02.05", "P0TC", "TRACTOCAMIONES", "TM2 SUR", "REPERCUSIÓN"])
    ws.append([datetime(2026, 2, 1), "VENTA", "TRACTOCAMIÓN X", "40BUP0TC001", 200, 0.5, 100, 1200, 120000, "T2 - UF1", "3701", "3701.02.05", "P0TC", "TRACTOCAMIONES", "TM2 SUR", "REPERCUSIÓN"])
    ws.append([datetime(2026, 2, 1), "VENTA", "REGULARIZACIONES", "40I10", None, None, None, None, 5000, "MAQUINARIA", "40I10", "40I10", "40I10", "REGULARIZACIONES", "REGULARIZACIONES", "STOCK REALIZABLE"])
    ws.append([datetime(2026, 3, 1), "PTTO VENTA", "TRACTOCAMIÓN X", "40BUP0TC001", 192, None, 192, 1200, 230400, "PRESUPUESTO", "PPTO", "PRESUPUESTO", "P0TC", "TRACTOCAMIONES", "PPTO", "PTTO VENTA"])
    eq = wb.create_sheet("LISTA DE EQUIPOS")
    eq.append(["CÓDIGO EQUIPO", "TARIFA HORA", "TIPO DE EQUIPO", "PLACA", "GPS", None, "PLACA"])
    eq.append(["40BUP0TC001", 1200, "TRACTOCAMIÓN X", None, None, None, "AAA001"])
    fam = wb.create_sheet("LISTA FAMILIAS")
    fam.append(["P0TC", "TRACTOCAMIONES"])
    pr = wb.create_sheet("LISTA DE PROYECTOS")
    pr.append(["No PROYECTO", "NOMBRE CORTO", "NOMBRE PROYECTO", "CONTROL DE GESTIÓN", "JEFE DE OBRA", "JEFE DE PRODUCCIÓN", "PROYECTO"])
    pr.append(["3701", "T2 - UF1", "T2 - UF1 - R4513", "A", "B", "C", "TM2 SUR"])
    cv = wb.create_sheet("COSTES VS VENTA")
    cv.append(["FECHA", "TIPO", "IMPORTE"])
    cv.append([datetime(2026, 1, 1), "VENTA INCURRIDA", 100000])
    cv.append([datetime(2026, 2, 1), "VENTA INCURRIDA", 125000])
    p = tmp_path / "master.xlsx"
    wb.save(p)
    return parse(p)


def test_parse(book):
    assert short_code("40BUP0TC001") == "P0TC001" and short_code("40I10") == "40I10"
    assert book.assets["40BUP0TC001"]["plate"] == "AAA001" and book.projects["3701"]["group_name"] == "TM2 SUR"
    kinds = [x.kind for x in book.lines]
    assert kinds == ["real", "real", "real", "presupuesto"]
    assert book.lines[1].hours == 200 and book.lines[1].pct == 0.5
    assert rate_history(book.lines) == [("40BUP0TC001", datetime(2026, 1, 1).date(), 1000), ("40BUP0TC001", datetime(2026, 2, 1).date(), 1200)]
    s = {x["mes"]: x for x in summary(book)}
    assert s["2026-02"]["real"] == 125000 and s["2026-02"]["cuadra"] and s["2026-03"]["presupuesto"] == 230400


@pytest.mark.skipif(not os.environ.get("FLEETMAP_TEST_DSN"), reason="sin FLEETMAP_TEST_DSN")
def test_postgres(book):
    psycopg = pytest.importorskip("psycopg")
    with psycopg.connect(os.environ["FLEETMAP_TEST_DSN"], autocommit=True) as conn:
        conn.execute("DROP SCHEMA IF EXISTS public CASCADE; CREATE SCHEMA public;")
        for f in ("schema.sql", "002_operacion_facturacion.sql", "003_repercusiones.sql"):
            conn.execute((DB / f).read_text(encoding="utf-8"))
        sql = to_sql(book)
        conn.execute(sql)
        conn.execute(sql)  # recargar reemplaza los meses, no duplica
        assert conn.execute("SELECT count(*), sum(amount) FROM repercussion").fetchone() == (4, 455400)
        assert float(conn.execute("SELECT asset_rate_at('40BUP0TC001', '2026-01-15')").fetchone()[0]) == 1000
        assert conn.execute("SELECT equipment_id FROM asset WHERE code = '40BUP0TC001'").fetchone()[0] == "P0TC001"
