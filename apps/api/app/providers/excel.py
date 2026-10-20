"""Proveedor basado en el libro "ARRANQUE DIARIO DE EQUIPOS" (40BU).

Hojas usadas:
- ARRANQUE   maestro operativo: código, descripción, placa, modelo, encargado y UF por día del corte.
- FORMATOS   estado diario del corte (T trabajando, D disponible, V varado) y tramo por día.
- HISTORICO  movimientos desde el inicio (fecha, código, estado, tramo, ubicación, observación):
             se arrastra el último estado conocido para reconstruir días anteriores al corte.
- FEYM511.00.CO  familia de cada equipo.

El libro no trae horas reales ni posición GPS: un día T cuenta el turno completo como trabajado,
D como standby y V como avería. Las posiciones salen de sites.json (corredor de cada proyecto) y
quedan marcadas como provisionales hasta tener coordenadas reales o telemetría.
"""
import asyncio
import hashlib
import json
import math
import re
import unicodedata
from collections.abc import AsyncIterator
from datetime import date, datetime, timedelta
from pathlib import Path

from ..domain import catalog as cat
from ..domain.models import DailyMetrics, Equipment, LiveStatus

LETTER = {"T": "op", "D": "standby", "V": "down", "VD": "down", "OTS": "off"}
WORD = {"TRABAJANDO": "op", "DISPONIBLE": "standby", "VARADO": "down", "FUERA DE SERVICIO": "off", "VENDIDO": "off"}
# Los puentes de cada zona se agrupan con su proyecto.
TRAMO_SITE = {"TM2 PUENTES SUR": "TM2 SUR", "TM2 PUENTES NORTE": "TM2 NORTE", "TM1 PS": "TM1 SUR", "TM1 PN": "TM1 NORTE"}
# Familia → icono disponible (el resto se dibuja como punto de estado).
FAMILY_ICON = {
    "BULLDOZER": "BUL", "CARGADORES": "CAR", "COMPACTADORES": "COM", "EXCAVADORAS": "EXC", "MOTONIVELADORAS": "MOT",
    "MINICARGADORES": "MIN", "VOLQUETAS DOBLETROQUE": "VOL", "TRACTOCAMIONES": "VOL", "CAMIONES": "VOL",
}


def _slug(s: str) -> str:
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]+", "-", s.upper()).strip("-")


def _day(v) -> date | None:
    return v.date() if isinstance(v, datetime) else v if isinstance(v, date) else None


def _uf_key(loc: str) -> str | None:
    """'TM1 UF05', 'T1 - UF5', 'T2 - UF0 REHABILITACIÓN' → '5' / '0'; puentes → 'PS' / 'PN'."""
    u = loc.upper()
    if m := re.search(r"UF\s*0*(\d+)", u):
        return m.group(1)
    if "PS" in u.split() or "PUENTES ZONA SUR" in u:
        return "PS"
    if "PN" in u.split() or "PUENTES ZONA NORTE" in u:
        return "PN"
    return None


def _h(s: str) -> float:
    return int(hashlib.md5(s.encode()).hexdigest()[:8], 16) / 0xFFFFFFFF


def _along(path: list[list[float]], f: float) -> tuple[float, float, float]:
    """Punto a la fracción f del recorrido [lat, lon] y su rumbo (para desplazar a un lado)."""
    segs = [math.dist(path[i], path[i + 1]) for i in range(len(path) - 1)]
    d, total = f * sum(segs), 0.0
    for i, s in enumerate(segs):
        if total + s >= d or i == len(segs) - 1:
            t = (d - total) / (s or 1)
            a, b = path[i], path[i + 1]
            return a[0] + (b[0] - a[0]) * t, a[1] + (b[1] - a[1]) * t, math.atan2(b[0] - a[0], b[1] - a[1])
        total += s
    return path[0][0], path[0][1], 0.0


class ExcelProvider:
    def __init__(self, path: str, sites_path: str, today: str, history_days: int):
        self.path, self.sites_path = Path(path), Path(sites_path)
        self.today, self.history_days = date.fromisoformat(today), history_days
        self._mtime = -1.0
        self._load()

    # --- carga -------------------------------------------------------------------------------
    def _reload_if_changed(self) -> None:
        try:
            m = self.path.stat().st_mtime + self.sites_path.stat().st_mtime
        except OSError:
            return
        if m != self._mtime:
            self._load()

    def _load(self) -> None:
        import openpyxl  # sólo se necesita con este proveedor

        self._mtime = self.path.stat().st_mtime + self.sites_path.stat().st_mtime
        cfg = json.loads(self.sites_path.read_text(encoding="utf-8"))
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)

        family: dict[str, str] = {}
        for r in wb["FEYM511.00.CO"].iter_rows(min_row=3, values_only=True):
            if r[1] and r[3]:
                family[str(r[1]).strip()] = str(r[3]).strip()

        rows = list(wb["ARRANQUE"].iter_rows(values_only=True))
        hdr = rows[6]
        day_cols = [(i, _day(v)) for i, v in enumerate(hdr) if _day(v)]
        master: dict[str, dict] = {}
        uf_by_day: dict[str, dict[date, str]] = {}
        for r in rows[7:]:
            code = str(r[2]).strip() if r[2] else ""
            if not code or code in master:
                continue
            master[code] = {"desc": r[4] or "", "plate": r[5] or "", "model": r[6] or "", "clase": r[7] or "", "manager": r[8] or ""}
            uf_by_day[code] = {d: str(r[i]).strip() for i, d in day_cols if isinstance(r[i], str) and r[i].strip()}

        rows = list(wb["FORMATOS"].iter_rows(values_only=True))
        hdr = rows[6]
        dates = [(i, _day(v)) for i, v in enumerate(hdr) if _day(v)]
        st_cols = [(i, d) for i, d in dates if i < 40]
        tr_cols = [(i, d) for i, d in dates if 40 < i < 76]
        state_by_day: dict[str, dict[date, str]] = {}
        tramo_by_day: dict[str, dict[date, str]] = {}
        for r in rows[7:]:
            code = str(r[2]).strip() if r[2] else ""
            if not code:
                continue
            state_by_day.setdefault(code, {}).update({d: LETTER[str(r[i]).strip()] for i, d in st_cols if str(r[i]).strip() in LETTER})
            tramo_by_day.setdefault(code, {}).update({d: str(r[i]).strip() for i, d in tr_cols if isinstance(r[i], str) and r[i].strip()})
            master.setdefault(code, {"desc": r[4] or "", "plate": r[5] or "", "model": "", "clase": r[7] or "", "manager": r[6] or ""})

        events: dict[str, list[tuple[date, str, str, str, str]]] = {}
        for r in wb["HISTORICO"].iter_rows(min_row=4, values_only=True):
            d, code = _day(r[1]), (str(r[2]).strip() if r[2] else "")
            if d and code and r[4] in WORD:
                events.setdefault(code, []).append((d, WORD[r[4]], str(r[5] or "").strip(), str(r[6] or "").strip(), str(r[8] or "").strip()))
        for ev in events.values():
            ev.sort(key=lambda e: e[0])
        wb.close()

        # --- obras: los grandes proyectos (y bodegas/talleres/plantas) de sites.json ----------
        self.sites = cfg["sites"]
        site_ids = {s["id"] for s in self.sites}
        fallback = cfg.get("fallbackSite", self.sites[0]["id"])

        def site_of(tramo: str) -> str:
            t = TRAMO_SITE.get(tramo, tramo)
            return _slug(t) if _slug(t) in site_ids else fallback

        self.start = self.today - timedelta(days=self.history_days)
        self.n_days = self.history_days + 1
        days = [self.start + timedelta(days=k) for k in range(self.n_days)]
        shift = {s["id"]: s for s in self.sites}

        fleet: list[Equipment] = []
        m = {k: [0.0] * 0 for k in ("work", "idle", "standby", "down", "fuel")}
        live: dict[str, LiveStatus] = {}
        type_names: dict[str, str] = {}
        now_ms = int(datetime.combine(self.today, datetime.min.time()).timestamp() * 1000)

        for code, info in sorted(master.items()):
            fam = family.get(code) or str(info["clase"] or "OTROS")
            type_id = _slug(fam)
            type_names[type_id] = fam.capitalize()
            ev = events.get(code, [])
            sday, tday = state_by_day.get(code, {}), tramo_by_day.get(code, {})

            # estado y tramo por día: FORMATOS manda en el corte; antes, último evento del histórico
            st_series, last_tramo, last_uf, last_note, k = [], "", "", "", 0
            cur_state, corte = "off", min(sday, default=self.today + timedelta(days=1))
            for d in days:
                while k < len(ev) and ev[k][0] <= d:
                    _, cur_state, tr, uf, note = ev[k]
                    last_tramo, last_uf, last_note = tr or last_tramo, uf or last_uf, note
                    k += 1
                # en el corte, un día sin letra es un día sin arranque reportado
                s = sday.get(d, "off") if d >= corte else cur_state
                if d in tday:
                    last_tramo = tday[d]
                if d in uf_by_day.get(code, {}):
                    last_uf = uf_by_day[code][d]
                st_series.append(s)
            sold = bool(ev) and ev[-1][1] == "off" and not any(v != "off" for v in sday.values())
            if sold or (not last_tramo and not ev and not sday):
                continue
            site_id = site_of(last_tramo)
            site = shift[site_id]
            for d, s in zip(days, st_series):
                hours = (site.get("satHours", 6) if d.weekday() == 5 else site.get("sunHours", 0) if d.weekday() == 6 else site.get("shiftHours", 10))
                m["work"].append(hours if s == "op" else 0.0)
                m["idle"].append(0.0)
                m["standby"].append(hours if s == "standby" else 0.0)
                m["down"].append(hours if s == "down" else 0.0)
                m["fuel"].append(0.0)

            lat, lon = self._place(site, code, last_uf)
            fleet.append(Equipment(
                id=code, pin=str(info["plate"] or ""), oem="", model=str(info["model"] or info["desc"]), type_id=type_id,
                site_id=site_id, client_id=site["client"], lat=lat, lon=lon, rate=0.0, standby_factor=cat.STANDBY_FACTOR,
                year=0, hour_meter=0.0, route=None, description=str(info["desc"]), plate=str(info["plate"] or ""),
                uf=last_uf or last_tramo, manager=str(info["manager"] or ""), icon=FAMILY_ICON.get(fam.upper()),
            ))
            # estado actual = último reporte (el día en curso suele estar a medio diligenciar)
            reported = [d for d in sday if d <= self.today]
            last_report = max(reported, default=ev[-1][0] if ev else self.start)
            state = sday[last_report] if reported else (ev[-1][1] if ev else "off")
            live[code] = LiveStatus(id=code, state=state, lat=lat, lon=lon,
                                    ts=now_ms - (self.today - last_report).days * 86_400_000,
                                    fault=last_note if state == "down" and last_note else None)

        self.fleet, self.m, self.live, self.days = fleet, m, live, days
        self.types = [{"id": t, "name": n, "rate": 0, "fuel": 0, "models": []} for t, n in sorted(type_names.items(), key=lambda x: x[1])]
        self.clients = cfg["clients"]

    def _place(self, site: dict, code: str, uf: str) -> tuple[float, float]:
        """Ubica el equipo sobre el corredor del proyecto; la UF fija el tramo dentro del corredor."""
        path = site.get("path")
        if not path:  # sitio puntual (planta, taller): disco alrededor del punto
            a, r = _h(code + "a") * math.tau, math.sqrt(_h(code + "r")) * site.get("spread", 0.01)
            return round(site["lat"] + math.sin(a) * r, 5), round(site["lon"] + math.cos(a) * r, 5)
        ufs = site.get("ufs") or []
        key = _uf_key(uf)
        if key in ufs:
            i = ufs.index(key)
            f = (i + 0.15 + 0.7 * _h(code)) / len(ufs)
        else:
            f = _h(code)
        lat, lon, heading = _along(path, f)
        side = (_h(code + "s") - 0.5) * 2 * site.get("spread", 0.004)
        return round(lat + math.cos(heading) * side, 5), round(lon - math.sin(heading) * side, 5)

    # --- contrato TelemetryProvider ----------------------------------------------------------
    def get_catalog(self) -> dict:
        self._reload_if_changed()
        lats = [p[0] for s in self.sites for p in (s.get("path") or [[s["lat"], s["lon"]]])]
        lons = [p[1] for s in self.sites for p in (s.get("path") or [[s["lat"], s["lon"]]])]
        sites = [{k: v for k, v in s.items() if k not in ("ufs",)} | {"sunday": s.get("sunHours", 0) > 0} for s in self.sites]
        return {
            "currency": "COP", "types": self.types,
            "states": [
                # sin ralentí: el libro no lo distingue
                {"id": "op", "name": "Trabajando", "color": "#12a87a"},
                {"id": "standby", "name": "Disponible", "color": "#2f86e0"}, {"id": "down", "name": "Varado", "color": "#e5384f"},
                {"id": "off", "name": "Sin reporte", "color": "#8a94a6"},
            ], "clients": self.clients, "sites": sites,
            "bbox": {"latMin": min(lats) - 0.1, "latMax": max(lats) + 0.1, "lonMin": min(lons) - 0.1, "lonMax": max(lons) + 0.1},
        }

    async def get_fleet(self) -> list[Equipment]:
        self._reload_if_changed()
        return self.fleet

    async def get_daily_metrics(self, date_from: str, date_to: str) -> DailyMetrics:
        self._reload_if_changed()
        a = max((date.fromisoformat(date_from) - self.start).days, 0)
        b = min((date.fromisoformat(date_to) - self.start).days, self.n_days - 1)
        if b < a:
            a = b = max(0, min(a, self.n_days - 1))
        out: dict[str, list[float]] = {}
        for k, src in self.m.items():
            col: list[float] = []
            for e in range(len(self.fleet)):
                base = e * self.n_days
                col.extend(src[base + a: base + b + 1])
            out[k] = col
        return DailyMetrics(days=[d.isoformat() for d in self.days[a: b + 1]], equipment_ids=[eq.id for eq in self.fleet], **out)

    async def get_live_status(self) -> list[LiveStatus]:
        self._reload_if_changed()
        return list(self.live.values())

    async def live_events(self) -> AsyncIterator[LiveStatus]:
        """Sin flujo en vivo: cuando el libro cambia en disco se recarga y se emiten los cambios de estado."""
        while True:
            await asyncio.sleep(30)
            before = {k: v.state for k, v in self.live.items()}
            self._reload_if_changed()
            for k, s in self.live.items():
                if before.get(k) not in (None, s.state):
                    yield s.model_copy(update={"prev": before[k]})
